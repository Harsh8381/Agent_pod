import io
import math
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from xml.sax.saxutils import escape

from html import escape as html_escape

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from frontend.claim_generation_page import register_generated_bill
from frontend.ui.components import badge, card_title, initials, page_title, render_banner, render_tiles
from frontend.ui.shell import render_page


DEFAULT_WORKBOOK = Path(__file__).resolve().parents[2] / "IPD_25_Patient_Billing_Demo.xlsx"
SHEET_ALIASES = {
    "IPD Patient & Services": ("IPD Patient & Services",),
    "IPD Charge Lines": ("IPD Charge Lines",),
    "CPT & ICD Code Reference": (
        "CPT & ICD Code Reference",
        "IPD_CPT_ICD_Code_Reference",
    ),
}
REQUIRED_COLUMNS = {
    "IPD Patient & Services": ("Encounter ID", "Patient ID", "Patient Name"),
    "IPD Charge Lines": (
        "Encounter ID", "Charge Line ID", "Service Date", "Service Description", "Revenue Code",
        "CPT / HCPCS", "ICD-10-CM Pointer", "Quantity", "Unit Charge ($)", "Line Charge ($)",
        "Payer Responsibility (%)",
    ),
    "CPT & ICD Code Reference": ("Code Type", "Code"),
}
IDENTIFIER_HEADER = re.compile(r"(?:^|\b)(?:id|code|cpt|hcpcs|icd|drg|revenue|pointer|ndc|member)(?:\b|$)", re.I)
CENT = Decimal("0.01")


def _cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return ""
        header = str(cell.parent.cell(1, cell.column).value or "")
        if IDENTIFIER_HEADER.search(header):
            number_format = str(cell.number_format or "").split(";")[0]
            if isinstance(value, (int, float)) and float(value).is_integer():
                width = len(number_format) if re.fullmatch(r"0+", number_format) else 0
                return f"{int(value):0{width}d}" if width else str(int(value))
        return format(value, ".15g") if isinstance(value, float) else str(value)
    return str(value).strip()


def _read_sheet(worksheet):
    rows = worksheet.iter_rows()
    header_cells = next(rows, ())
    headers = [str(cell.value or "").strip() for cell in header_cells]
    if not headers or not any(headers):
        return pd.DataFrame()
    records = []
    for row in rows:
        values = [_cell_text(cell) for cell in row[:len(headers)]]
        if any(value != "" for value in values):
            records.append(dict(zip(headers, values)))
    return pd.DataFrame(records, columns=headers)


@st.cache_data(show_spinner=False, ttl="10m", max_entries=2)
def load_ipd_workbook(file_path, modified_time_ns):
    del modified_time_ns
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"IPD billing workbook was not found: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as error:
        raise ValueError(f"The IPD billing workbook could not be opened: {error}") from error

    loaded = {}
    try:
        sheet_lookup = {re.sub(r"[^a-z0-9]", "", name.casefold()): name for name in workbook.sheetnames}
        for requested_name, aliases in SHEET_ALIASES.items():
            actual_name = next((name for name in aliases if name in workbook.sheetnames), None)
            if actual_name is None:
                actual_name = next(
                    (sheet_lookup.get(re.sub(r"[^a-z0-9]", "", name.casefold())) for name in aliases
                     if sheet_lookup.get(re.sub(r"[^a-z0-9]", "", name.casefold()))),
                    None,
                )
            if actual_name is None:
                raise ValueError(
                    f"Required sheet '{requested_name}' is missing. Available sheets: "
                    + ", ".join(workbook.sheetnames)
                )
            frame = _read_sheet(workbook[actual_name])
            missing = [column for column in REQUIRED_COLUMNS[requested_name] if column not in frame.columns]
            if missing:
                raise ValueError(f"Sheet '{actual_name}' is missing required columns: {', '.join(missing)}")
            loaded[requested_name] = frame.fillna("")
    finally:
        workbook.close()
    return loaded


def _text(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _split_codes(value):
    return [part.strip() for part in re.split(r"[;,\n]+", _text(value)) if part.strip()]


def _code_key(value):
    return re.sub(r"\s+", "", _text(value)).upper()


def _decimal(value):
    text = _text(value).replace("$", "").replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return number if number.is_finite() else None


def _date_value(value):
    text = _text(value)
    if not text:
        return None
    parsed = pd.to_datetime(text, errors="coerce")
    return None if pd.isna(parsed) else parsed.date()


def _display_date(value, time_value=None):
    """'2026-01-04T00:00:00' + '08:30' -> '04 Jan 2026 08:30' for display."""
    parsed = _date_value(value)
    text = parsed.strftime("%d %b %Y") if parsed else _text(value)
    time_text = _text(time_value)
    return f"{text} {time_text}".strip() or "—"


def _money(value):
    return f"${Decimal(value).quantize(CENT, rounding=ROUND_HALF_UP):,.2f}"


def _build_reference_map(reference):
    result = {}
    for row in reference.to_dict("records"):
        key = _code_key(row.get("Code"))
        if key:
            result.setdefault(key, row)
    return result


def _make_charge_lines(charges):
    lines = []
    for row in charges.to_dict("records"):
        quantity = _decimal(row.get("Quantity"))
        unit_charge = _decimal(row.get("Unit Charge ($)"))
        safe_quantity = quantity if quantity is not None else Decimal("0")
        safe_unit_charge = unit_charge if unit_charge is not None else Decimal("0")
        line_charge = (safe_quantity * safe_unit_charge).quantize(CENT, rounding=ROUND_HALF_UP)
        lines.append({
            **row,
            "Quantity": safe_quantity,
            "Unit Charge ($)": safe_unit_charge.quantize(CENT, rounding=ROUND_HALF_UP),
            "Line Charge ($)": line_charge,
        })
    return lines


def _collect_codes(patient, charges):
    icd_codes = []
    cpt_codes = []
    for field in ("Primary ICD-10-CM", "Secondary ICD-10-CM"):
        icd_codes.extend(_split_codes(patient.get(field)))
    cpt_codes.extend(_split_codes(patient.get("Key CPT / HCPCS")))
    for row in charges.to_dict("records"):
        icd_codes.extend(_split_codes(row.get("ICD-10-CM Pointer")))
        cpt_codes.extend(_split_codes(row.get("CPT / HCPCS")))
    return list(dict.fromkeys(icd_codes)), list(dict.fromkeys(cpt_codes))


def _code_details(patient, charges, references):
    icd_codes, cpt_codes = _collect_codes(patient, charges)
    charge_services = {}
    for row in charges.to_dict("records"):
        for code in _split_codes(row.get("CPT / HCPCS")) + _split_codes(row.get("ICD-10-CM Pointer")):
            service = _text(row.get("Service Description"))
            if service:
                charge_services.setdefault(_code_key(code), []).append(service)
    rows = []
    for code_type, codes in (("ICD-10-CM", icd_codes), ("CPT/HCPCS", cpt_codes)):
        for code in codes:
            reference = references.get(_code_key(code), {})
            rows.append({
                "Code type": code_type,
                "Code": code,
                "Disease / procedure": _text(reference.get("Plain-Language Diagnosis / Procedure")) or "Not found in code reference",
                "Clinical category": _text(reference.get("Clinical Category")),
                "Treatment / use": _text(reference.get("Corresponding Disease / Treatment / Use")),
                "Encounter service details": "; ".join(dict.fromkeys(charge_services.get(_code_key(code), []))),
                "Reference note": _text(reference.get("Important Note")),
            })
    return rows, icd_codes, cpt_codes


def _validate_encounter(patient, charges, references, lines):
    issues = []
    required_values = ("Patient ID", "Patient Name", "Payer", "Admission Date", "Discharge Date", "Primary Diagnosis", "MS-DRG")
    missing_values = [field for field in required_values if not _text(patient.get(field))]
    if missing_values:
        issues.append("Missing patient/encounter data: " + ", ".join(missing_values))

    admission = _date_value(patient.get("Admission Date"))
    discharge = _date_value(patient.get("Discharge Date"))
    if not admission:
        issues.append("Admission date is missing or invalid.")
    if not discharge:
        issues.append("Discharge date is missing or invalid.")
    if admission and discharge:
        if discharge < admission:
            issues.append("Discharge date occurs before admission date.")
        else:
            los = _decimal(patient.get("Length of Stay (Days)"))
            expected_los = Decimal(str((discharge - admission).days))
            if los is None or los < 0 or los != los.to_integral_value():
                issues.append("Length of stay is missing or is not a non-negative whole number.")
            elif los != expected_los:
                issues.append(f"Length of stay is {los} day(s), but the admission/discharge dates span {expected_los} day(s).")

    line_ids = [_text(row.get("Charge Line ID")) for row in charges.to_dict("records")]
    seen_line_ids = set()
    duplicate_line_ids = set()
    for line_id in line_ids:
        if line_id and line_id in seen_line_ids:
            duplicate_line_ids.add(line_id)
        seen_line_ids.add(line_id)
    if duplicate_line_ids:
        issues.append("Duplicate charge line IDs: " + ", ".join(sorted(duplicate_line_ids)))
    duplicate_fields = ("Service Date", "Service Description", "Revenue Code", "CPT / HCPCS", "Quantity", "Unit Charge ($)")
    if not charges.empty:
        duplicate_mask = charges.duplicated(subset=list(duplicate_fields), keep=False)
        duplicate_count = int(duplicate_mask.sum())
        if duplicate_count:
            issues.append(f"Potential duplicate charges: {duplicate_count} lines repeat the same date, service, code, quantity and unit charge.")

    for index, row in enumerate(charges.to_dict("records")):
        quantity = _decimal(row.get("Quantity"))
        unit_charge = _decimal(row.get("Unit Charge ($)"))
        if quantity is None or quantity <= 0:
            issues.append(f"Charge line {_text(row.get('Charge Line ID')) or index + 1} has missing or invalid quantity.")
        if unit_charge is None or unit_charge < 0:
            issues.append(f"Charge line {_text(row.get('Charge Line ID')) or index + 1} has missing or invalid unit charge.")
        recorded = _decimal(row.get("Line Charge ($)"))
        recalculated = lines[index]["Line Charge ($)"]
        if recorded is None or recorded.quantize(CENT, rounding=ROUND_HALF_UP) != recalculated:
            issues.append(f"Charge line {_text(row.get('Charge Line ID')) or index + 1} total does not equal quantity × unit charge.")
    reference_keys = set(references)
    _, icd_codes, cpt_codes = _code_details(patient, charges, references)
    unknown_codes = [code for code in icd_codes + cpt_codes if _code_key(code) not in reference_keys]
    if unknown_codes:
        issues.append("Codes not found in the workbook reference: " + ", ".join(dict.fromkeys(unknown_codes)))

    actual_total = sum((line["Line Charge ($)"] for line in lines), Decimal("0.00")).quantize(CENT)
    stated_total = _decimal(patient.get("Estimated Total Charges ($)"))
    if stated_total is not None and stated_total.quantize(CENT, rounding=ROUND_HALF_UP) != actual_total:
        issues.append(f"Encounter total is {_money(stated_total)}, while recalculated charge lines total {_money(actual_total)}.")
    return issues


def _payer_share(row):
    value = _decimal(row.get("Payer Responsibility (%)"))
    if value is None:
        return Decimal("0")
    if value > 1:
        value /= Decimal("100")
    return min(max(value, Decimal("0")), Decimal("1"))


def _build_pdf(patient, lines, gross, payer_amount, patient_amount):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=landscape(letter), leftMargin=0.45 * inch, rightMargin=0.45 * inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallBill", parent=styles["BodyText"], fontSize=7.5, leading=9))
    styles.add(ParagraphStyle(name="BillRight", parent=styles["SmallBill"], alignment=TA_RIGHT))
    story = [
        Paragraph("IPD itemized bill", styles["Title"]),
        Paragraph("Demonstration estimate only. Not a finalized claim or determination of benefits.", styles["BodyText"]),
        Spacer(1, 10),
        Paragraph(
            "<b>Encounter:</b> " + escape(_text(patient.get("Encounter ID")))
            + " &nbsp;&nbsp; <b>Patient:</b> " + escape(_text(patient.get("Patient Name")))
            + " &nbsp;&nbsp; <b>Payer:</b> " + escape(_text(patient.get("Payer"))),
            styles["BodyText"],
        ),
        Paragraph(
            "<b>Admission:</b> " + escape(_text(patient.get("Admission Date")))
            + " &nbsp;&nbsp; <b>Discharge:</b> " + escape(_text(patient.get("Discharge Date")))
            + " &nbsp;&nbsp; <b>DRG:</b> " + escape(_text(patient.get("MS-DRG"))),
            styles["BodyText"],
        ),
        Spacer(1, 10),
    ]
    table_data = [["Service date / description", "Revenue", "CPT/HCPCS", "Qty", "Unit charge", "Line charge"]]
    for line in lines:
        description = escape(_text(line.get("Service Description")))
        department = escape(_text(line.get("Department / Cost Center")))
        detail = "<b>" + description + "</b>"
        if department:
            detail += "<br/>" + department
        table_data.append([
            Paragraph(detail, styles["SmallBill"]),
            escape(_text(line.get("Revenue Code"))),
            escape(_text(line.get("CPT / HCPCS"))),
            str(line["Quantity"]),
            _money(line["Unit Charge ($)"]),
            _money(line["Line Charge ($)"]),
        ])
    table = Table(table_data, repeatRows=1, colWidths=[3.25 * inch, 0.7 * inch, 0.9 * inch, 0.45 * inch, 0.9 * inch, 0.9 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163b3a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5d1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f6f4")]),
        ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.extend([table, Spacer(1, 12)])
    totals = Table([
        ["Gross total", _money(gross)],
        ["Estimated payer amount", _money(payer_amount)],
        ["Estimated patient amount", _money(patient_amount)],
    ], colWidths=[2.2 * inch, 1.4 * inch], hAlign="RIGHT")
    totals.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.HexColor("#163b3a")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(totals)
    document.build(story)
    return buffer.getvalue()


BILL_COLUMNS = (
    "Charge Line ID", "Service Date", "Service Category", "Service Description", "Revenue Code",
    "CPT / HCPCS", "ICD-10-CM Pointer", "Quantity", "Unit Charge ($)", "Line Charge ($)",
)


def _charge_table(frame, columns):
    table = frame[columns].copy()
    if "Service Date" in table.columns:
        table["Service Date"] = table["Service Date"].map(_display_date)
    st.dataframe(table, hide_index=True, width="stretch")


def _patient_card(patient):
    name = _text(patient.get("Patient Name")) or "Unnamed patient"
    details = (
        ("Admission", _display_date(patient.get("Admission Date"), patient.get("Admission Time"))),
        ("Discharge", _display_date(patient.get("Discharge Date"), patient.get("Discharge Time"))),
        ("Ward / bed", f"{_text(patient.get('Unit / Ward'))} · {_text(patient.get('Room / Bed'))}"),
        ("Length of stay", f"{_text(patient.get('Length of Stay (Days)')) or 'N/A'} days"),
        ("Payer", _text(patient.get("Payer")) or "Not recorded"),
        ("MS-DRG", f"{_text(patient.get('MS-DRG'))} · {_text(patient.get('MS-DRG Description'))}"),
    )
    st.markdown(
        '<div class="nc-banner-card">'
        f'<div class="nc-patient"><div class="nc-initials lg">{html_escape(initials(name))}</div>'
        f'<div><b style="font-size:.95rem">{html_escape(name)}</b><small>Patient {html_escape(_text(patient.get("Patient ID")))} · '
        f'{html_escape(_text(patient.get("Gender")))} · DOB {html_escape(_display_date(patient.get("DOB")))}</small></div></div>'
        '<div class="nc-badge-row">'
        + badge(f'Primary {_text(patient.get("Primary ICD-10-CM"))} · {_text(patient.get("Primary Diagnosis"))}', "violet")
        + (badge(f'Secondary {_text(patient.get("Secondary ICD-10-CM"))}', "gray") if _text(patient.get("Secondary ICD-10-CM")) else "")
        + "</div></div>",
        unsafe_allow_html=True,
    )
    render_tiles(details)


def _generate_bill(selected_id, patient, lines):
    gross = sum((line["Line Charge ($)"] for line in lines), Decimal("0.00")).quantize(CENT)
    payer_amount = sum(
        (line["Line Charge ($)"] * _payer_share(line) for line in lines), Decimal("0.00")
    ).quantize(CENT, rounding=ROUND_HALF_UP)
    patient_amount = (gross - payer_amount).quantize(CENT, rounding=ROUND_HALF_UP)
    st.session_state.ipd_generated_bill = {
        "encounter_id": selected_id,
        "patient": patient,
        "lines": lines,
        "gross": gross,
        "payer_amount": payer_amount,
        "patient_amount": patient_amount,
    }
    register_generated_bill(st.session_state.ipd_generated_bill, source="IPD")


def show_ipd_billing(workbook_path=DEFAULT_WORKBOOK):
    render_page("ipd_billing")
    page_title("Billing & RCM", "IPD Billing", "Inpatient encounter details, code validation and itemized charge review")
    path = Path(workbook_path)
    if not path.is_file():
        render_banner("IPD billing workbook not found", str(path), tone="danger", icon="&#9888;")
        return
    try:
        workbook = load_ipd_workbook(str(path), path.stat().st_mtime_ns)
    except (OSError, ValueError, KeyError) as error:
        render_banner("The IPD billing workbook could not be read", str(error), tone="danger", icon="&#9888;")
        return

    patients = workbook["IPD Patient & Services"]
    if patients.empty:
        st.info("No inpatient encounters are available in the workbook.")
        return
    patients = patients.copy()
    patients["Encounter ID"] = patients["Encounter ID"].map(_text)
    encounter_ids = list(dict.fromkeys(value for value in patients["Encounter ID"] if value))
    if not encounter_ids:
        render_banner("The patient sheet has no Encounter IDs to select.", tone="danger", icon="&#9888;")
        return
    names = {row["Encounter ID"]: _text(row.get("Patient Name")) for row in patients.to_dict("records")}
    requested_encounter_id = st.session_state.get("ipd_billing_encounter")
    encounter_index = encounter_ids.index(requested_encounter_id) if requested_encounter_id in encounter_ids else 0
    selected_id = st.selectbox(
        "Inpatient encounter", encounter_ids, index=encounter_index, key="ipd_billing_encounter",
        format_func=lambda value: f"{value} · {names.get(value) or 'Unnamed patient'}",
    )
    matching_patients = patients[patients["Encounter ID"] == selected_id]
    patient = matching_patients.iloc[0].to_dict()
    if len(matching_patients) > 1:
        st.warning(f"Encounter ID {selected_id} appears more than once; displaying the first patient row.")
    charges = workbook["IPD Charge Lines"]
    charges = charges[charges["Encounter ID"].map(_text) == selected_id].copy()
    references = _build_reference_map(workbook["CPT & ICD Code Reference"])
    lines = _make_charge_lines(charges)
    validation = _validate_encounter(patient, charges, references, lines)

    _patient_card(patient)
    if validation:
        render_banner(f"{len(validation)} validation item(s) need review", " · ".join(validation), tone="warn", icon="&#9888;")
    else:
        render_banner("Encounter, code, date and charge-total checks passed", tone="ok", icon="&#10003;")

    with st.container(border=True):
        code_rows, _, _ = _code_details(patient, charges, references)
        card_title("Encounter codes and services", "ICD-10-CM diagnoses and CPT/HCPCS services with reference validation")
        if code_rows:
            st.dataframe(pd.DataFrame(code_rows), hide_index=True, width="stretch")
        else:
            st.info("No ICD or CPT/HCPCS codes are recorded for this encounter.")

    display_lines = pd.DataFrame(lines) if lines else pd.DataFrame(columns=charges.columns)
    available_columns = [column for column in BILL_COLUMNS if column in display_lines.columns]
    with st.container(border=True):
        card_title("Itemized charge lines", f"{len(lines)} line(s) linked to {selected_id}", trailing=badge(f"{len(lines)} lines", "blue"))
        _charge_table(display_lines, available_columns)
        if not lines:
            render_banner("No charge lines are linked to this encounter; a bill cannot be generated.", tone="warn", icon="&#9888;")
            return
        st.button("Generate bill", key="ipd_generate_bill", type="primary", icon=":material/receipt_long:",
                  on_click=_generate_bill, args=(selected_id, patient, lines))

    generated = st.session_state.get("ipd_generated_bill")
    if not (generated and generated.get("encounter_id") == selected_id):
        return
    with st.container(border=True):
        card_title("Generated bill", "Payer and patient estimates from the workbook's payer-responsibility percentages",
                   trailing=badge("Draft", "amber"))
        render_tiles((
            ("Gross total", _money(generated["gross"])),
            ("Estimated payer amount", _money(generated["payer_amount"])),
            ("Estimated patient amount", _money(generated["patient_amount"])),
        ))
        _charge_table(pd.DataFrame(generated["lines"]), available_columns)
        st.caption("Estimates are not benefits determinations or finalized patient responsibility.")
        patient_slug = re.sub(r"[^A-Za-z0-9_-]+", "_", _text(patient.get("Patient Name"))).strip("_") or "patient"
        try:
            pdf_bytes = _build_pdf(
                generated["patient"], generated["lines"], generated["gross"],
                generated["payer_amount"], generated["patient_amount"],
            )
        except (ImportError, ValueError, OSError) as error:
            render_banner("PDF generation is unavailable", str(error), tone="danger", icon="&#9888;")
        else:
            st.download_button(
                "Download bill PDF", data=pdf_bytes, file_name=f"IPD_Bill_{patient_slug}_{selected_id}.pdf",
                mime="application/pdf", key="ipd_bill_pdf", icon=":material/download:",
            )