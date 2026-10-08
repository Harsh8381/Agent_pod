from __future__ import annotations

import json
import os
import re
import tempfile
import threading
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = PROJECT_ROOT / "Epic_Inspired_USA_50_Patient_Demo.xlsx"
JSON_PATH = PROJECT_ROOT / "Epic_Inspired_USA_50_Patient_Demo_Athena_Care_Gap.json"
BILLING_NOTICE = (
    "This draft contains gross chargemaster-based demo charges only. It is not a finalized claim, "
    "payer-allowed amount, or patient responsibility amount. Professional coding and billing validation is required."
)
TABLE_SHEETS = {
    "Patient_Master",
    "Encounters",
    "Problems",
    "Allergies",
    "Medications",
    "Vitals",
    "Labs",
    "Orders_Procedures",
    "Immunizations",
    "Appointments",
    "Clinical_Notes",
    "Insurance",
    "Billing_Claims",
    "Consent_Directives",
    "Care_Team",
    "Family_Social_Hx",
    "Data_Dictionary",
    "Sources",
}
_STORE_LOCK = threading.RLock()


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float):
        return value if value == value and abs(value) != float("inf") else None
    if hasattr(value, "item"):
        try:
            return _json_value(value.item())
        except (ValueError, TypeError):
            return None
    return str(value)


def _atomic_write(payload: dict[str, Any]) -> None:
    JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_path = tempfile.mkstemp(
        prefix=f".{JSON_PATH.stem}.", suffix=".tmp", dir=JSON_PATH.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary_path, JSON_PATH)
    finally:
        if os.path.exists(temporary_path):
            os.unlink(temporary_path)


def _workbook_payload() -> dict[str, Any]:
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"EHR source workbook not found: {WORKBOOK_PATH}")
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=False)
    sheets: dict[str, Any] = {}
    try:
        for worksheet in workbook.worksheets:
            values = [[_json_value(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
            if worksheet.title in TABLE_SHEETS:
                headers = [str(value).strip() if value is not None else f"Column_{index + 1}" for index, value in enumerate(values[0])]
                records = []
                for values_row in values[1:]:
                    if not any(value is not None for value in values_row):
                        continue
                    padded = values_row + [None] * max(0, len(headers) - len(values_row))
                    records.append(dict(zip(headers, padded)))
                sheets[worksheet.title] = {"headers": headers, "records": records}
            else:
                raw_rows = []
                for row in values:
                    while row and row[-1] is None:
                        row.pop()
                    raw_rows.append(row)
                sheets[worksheet.title] = {"raw_rows": raw_rows}
    finally:
        workbook.close()
    return {
        "format": "nuucare-workbook-json-v1",
        "source_workbook": WORKBOOK_PATH.name,
        "created_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "sheets": sheets,
    }


def ensure_json_store() -> dict[str, Any]:
    with _STORE_LOCK:
        if JSON_PATH.is_file():
            try:
                payload = json.loads(JSON_PATH.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as error:
                raise RuntimeError(f"EHR JSON store is invalid and was not overwritten: {error}") from error
            patient_sheet = payload.get("sheets", {}).get("Patient_Master") if isinstance(payload, dict) else None
            if (
                not isinstance(payload, dict)
                or not isinstance(payload.get("sheets"), dict)
                or not isinstance(patient_sheet, dict)
                or not isinstance(patient_sheet.get("records"), list)
            ):
                raise RuntimeError("EHR JSON store must contain a workbook-shaped 'sheets' object.")
            return payload
        payload = _workbook_payload()
        _atomic_write(payload)
        return payload


def _read_store() -> dict[str, Any]:
    return ensure_json_store()


def _write_store(payload: dict[str, Any]) -> None:
    with _STORE_LOCK:
        _atomic_write(payload)


def _sheet_records(payload: dict[str, Any], sheet_name: str, *, create: bool = False) -> list[dict[str, Any]]:
    sheets = payload["sheets"]
    if sheet_name not in sheets:
        if not create:
            raise KeyError(f"EHR sheet not found: {sheet_name}")
        sheets[sheet_name] = {"headers": [], "records": []}
    sheet = sheets[sheet_name]
    if "records" not in sheet:
        if not create:
            raise ValueError(f"EHR sheet is not a record table: {sheet_name}")
        sheet.clear()
        sheet.update({"headers": [], "records": []})
    return sheet["records"]


def _get_value(record: dict[str, Any], *names: str) -> Any:
    lowered = {str(key).casefold(): value for key, value in record.items()}
    return next((lowered[name.casefold()] for name in names if name.casefold() in lowered), None)


def _patient_rows(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return _sheet_records(payload, "Patient_Master")


def _find_patient(payload: dict[str, Any], patient_id: str) -> dict[str, Any]:
    for patient in _patient_rows(payload):
        if str(_get_value(patient, "Patient_ID", "patient_id", "MRN")) == str(patient_id):
            return patient
    raise KeyError(f"Patient not found: {patient_id}")


def _new_patient_id(records: list[dict[str, Any]]) -> str:
    numbers = []
    for record in records:
        value = str(_get_value(record, "Patient_ID") or "")
        suffix = "".join(character for character in value if character.isdigit())
        if suffix:
            numbers.append(int(suffix))
    return f"P{max(numbers, default=100000) + 1:06d}"


def _row_for_headers(headers: list[str], supplied: dict[str, Any]) -> dict[str, Any]:
    values = {str(key).casefold(): _json_value(value) for key, value in supplied.items()}
    row = {header: values.get(header.casefold()) for header in headers}
    for key, value in supplied.items():
        if key not in row:
            row[key] = _json_value(value)
    return row


def list_patients() -> list[dict[str, Any]]:
    payload = _read_store()
    patients = []
    for row in _patient_rows(payload):
        patients.append({
            **row,
            "patient_id": _get_value(row, "Patient_ID", "patient_id"),
            "display_name": _get_value(row, "Legal_Name") or " ".join(
                str(part) for part in (_get_value(row, "First_Name"), _get_value(row, "Last_Name")) if part
            ),
        })
    return patients


def get_patient_by_id(patient_id: str) -> dict[str, Any]:
    return dict(_find_patient(_read_store(), patient_id))


def get_patient_chart(patient_id: str) -> dict[str, Any]:
    payload = _read_store()
    patient = _find_patient(payload, patient_id)
    records: dict[str, list[dict[str, Any]]] = {}
    for sheet_name, sheet in payload["sheets"].items():
        rows = sheet.get("records")
        if sheet_name == "Patient_Master" or not isinstance(rows, list):
            continue
        related = [row for row in rows if str(_get_value(row, "Patient_ID", "patient_id")) == str(patient_id)]
        if related:
            records[sheet_name] = related
    return {"patient": dict(patient), "records": records, "record_counts": {name: len(rows) for name, rows in records.items()}}


def create_patient(payload: dict[str, Any]) -> dict[str, Any]:
    with _STORE_LOCK:
        store = _read_store()
        patients = _patient_rows(store)
        values = dict(payload or {})
        patient_id = values.get("Patient_ID") or values.get("patient_id") or _new_patient_id(patients)
        if any(str(_get_value(row, "Patient_ID")) == str(patient_id) for row in patients):
            raise ValueError(f"Patient already exists: {patient_id}")
        values.setdefault("Patient_ID", str(patient_id))
        if not values.get("MRN") and not values.get("mrn"):
            values["MRN"] = f"MRN{str(patient_id).lstrip('P')}"
        values.setdefault("Record_Status", "Active")
        sheet = store["sheets"]["Patient_Master"]
        row = _row_for_headers(sheet["headers"], values)
        sheet["records"].append(row)
        _write_store(store)
    return dict(row)


def update_patient(patient_id: str, updates: dict[str, Any]) -> dict[str, Any]:
    with _STORE_LOCK:
        store = _read_store()
        patient = _find_patient(store, patient_id)
        headers_by_fold = {str(key).casefold(): key for key in patient}
        for key, value in (updates or {}).items():
            actual_key = headers_by_fold.get(str(key).casefold(), key)
            patient[actual_key] = _json_value(value)
        _write_store(store)
    return dict(patient)


def delete_patient(patient_id: str) -> bool:
    with _STORE_LOCK:
        store = _read_store()
        patient = _find_patient(store, patient_id)
        store["sheets"]["Patient_Master"]["records"].remove(patient)
        for sheet_name, sheet in store["sheets"].items():
            rows = sheet.get("records")
            if isinstance(rows, list) and sheet_name != "Patient_Master":
                sheet["records"] = [row for row in rows if str(_get_value(row, "Patient_ID", "patient_id")) != str(patient_id)]
        for generated_sheet in ("Generated_Bills", "Generated_Bill_Lines"):
            generated = store["sheets"].get(generated_sheet)
            if generated is not None and not generated.get("records"):
                store["sheets"].pop(generated_sheet)
        _write_store(store)
    return True


def list_patient_records(patient_id: str, record_type: str | None = None) -> list[dict[str, Any]]:
    chart = get_patient_chart(patient_id)
    if record_type:
        return chart["records"].get(record_type, [])
    return [row for records in chart["records"].values() for row in records]


def replace_patient_records(patient_id: str, sheet_name: str, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    with _STORE_LOCK:
        store = _read_store()
        _find_patient(store, patient_id)
        rows = _sheet_records(store, sheet_name)
        headers = store["sheets"][sheet_name]["headers"]
        replacements = []
        for supplied in records or []:
            record = _row_for_headers(headers, supplied)
            record["Patient_ID"] = str(patient_id)
            replacements.append(record)
        others = [row for row in rows if str(_get_value(row, "Patient_ID", "patient_id")) != str(patient_id)]
        store["sheets"][sheet_name]["records"] = others + replacements
        _write_store(store)
    return replacements


def add_patient_record(patient_id: str, sheet_name: str, record: dict[str, Any]) -> dict[str, Any]:
    with _STORE_LOCK:
        store = _read_store()
        _find_patient(store, patient_id)
        rows = _sheet_records(store, sheet_name)
        headers = store["sheets"][sheet_name]["headers"]
        supplied = dict(record or {})
        supplied.setdefault("Patient_ID", str(patient_id))
        inserted = _row_for_headers(headers, supplied)
        rows.append(inserted)
        _write_store(store)
    return inserted


def export_patient_data(patient_id: str | None = None) -> dict[str, Any]:
    store = _read_store()
    if patient_id is None:
        return store
    chart = get_patient_chart(patient_id)
    sheets = {"Patient_Master": {"headers": store["sheets"]["Patient_Master"]["headers"], "records": [chart["patient"]]}}
    for name, records in chart["records"].items():
        sheets[name] = {"headers": store["sheets"][name]["headers"], "records": records}
    return {"format": store.get("format"), "sheets": sheets, "patient_count": 1}


def list_bills(patient_id: str | None = None) -> list[dict[str, Any]]:
    store = _read_store()
    bills = _sheet_records(store, "Generated_Bills") if "Generated_Bills" in store["sheets"] else []
    lines = _sheet_records(store, "Generated_Bill_Lines") if "Generated_Bill_Lines" in store["sheets"] else []
    filtered = [row for row in bills if patient_id is None or str(_get_value(row, "Patient_ID")) == str(patient_id)]
    for bill in filtered:
        bill_id = str(_get_value(bill, "Bill_ID"))
        bill["lines"] = [line for line in lines if str(_get_value(line, "Bill_ID")) == bill_id]
    return filtered


def _append_generated_row(store: dict[str, Any], sheet_name: str, headers: list[str], row: dict[str, Any]) -> dict[str, Any]:
    sheet = store["sheets"].setdefault(sheet_name, {"headers": headers, "records": []})
    if "records" not in sheet:
        sheet.clear()
        sheet.update({"headers": headers, "records": []})
    for header in headers:
        if header not in sheet["headers"]:
            sheet["headers"].append(header)
    normalized = _row_for_headers(sheet["headers"], row)
    sheet["records"].append(normalized)
    return normalized


def generate_patient_bill(patient_id: str, bill_request: dict[str, Any]) -> dict[str, Any]:
    from cpt_coder import classify_procedure_status, is_explicitly_performed
    from frontend.chargemaster import price_cpt_bill_lines

    request = dict(bill_request or {})
    services = request.get("services") or request.get("bill_lines") or []
    reviewed_services = request.get("reviewed_services") or []
    evidence = str(request.get("documentation") or request.get("evidence") or "")
    service_by_code: dict[str, dict[str, Any]] = {}
    for item in services:
        if not isinstance(item, dict):
            continue
        code = str(item.get("CPT/HCPCS Code") or item.get("CPT_HCPCS") or item.get("code") or "").strip()
        normalized_code = re.sub(r"[\s.-]", "", code).upper()
        if code and normalized_code:
            service = dict(item)
            service.pop("performed_confirmed", None)
            service_by_code[normalized_code] = service
    for reviewed in reviewed_services:
        if not isinstance(reviewed, dict):
            continue
        code = str(reviewed.get("code") or reviewed.get("CPT/HCPCS Code") or reviewed.get("CPT_HCPCS") or "").strip()
        normalized_code = re.sub(r"[\s.-]", "", code).upper()
        if not code or not normalized_code or reviewed.get("performed_confirmed") is not True:
            continue
        source = str(reviewed.get("source") or "consultation_review")
        procedure = str(reviewed.get("procedure") or reviewed.get("Extracted Procedure") or code)
        source_parts = set(source.split("+"))
        is_ehr_encounter_source = "ehr_encounter" in source_parts
        if not is_ehr_encounter_source:
            status = classify_procedure_status(
                re.sub(r"\bwithout\s+contrast\b", "", procedure, flags=re.IGNORECASE),
                re.sub(r"\bwithout\s+contrast\b", "", evidence, flags=re.IGNORECASE),
            )
            if status not in {"performed", "unknown"}:
                continue
        service = service_by_code.get(normalized_code, {})
        service.update({
            "CPT/HCPCS Code": code,
            "Extracted Procedure": str(reviewed.get("procedure") or reviewed.get("Extracted Procedure") or service.get("Extracted Procedure") or code),
            "Matched Procedure/Service": str(reviewed.get("description") or reviewed.get("Matched Procedure/Service") or service.get("Matched Procedure/Service") or code),
            "quantity": reviewed.get("quantity", service.get("quantity", 1)),
            "billing_source": source,
            "performed_confirmed": True,
        })
        service_by_code[normalized_code] = service

    performed = [
        item for item in services
        if isinstance(item, dict)
        and isinstance(item.get("Extracted Procedure"), str)
        and is_explicitly_performed(
            re.sub(r"\bwithout\s+contrast\b", "", item["Extracted Procedure"], flags=re.IGNORECASE),
            re.sub(r"\bwithout\s+contrast\b", "", evidence, flags=re.IGNORECASE),
        )
    ]
    performed_codes = {
        re.sub(r"[\s.-]", "", str(item.get("CPT/HCPCS Code") or item.get("CPT_HCPCS") or "")).upper()
        for item in performed
    }
    performed.extend(
        item for normalized_code, item in service_by_code.items()
        if normalized_code not in performed_codes and item.get("performed_confirmed") is True
    )
    if not performed:
        raise ValueError("No explicitly performed and verifiable CPT/HCPCS procedures were supplied; no draft bill was created.")

    deduplicated_performed = {}
    for item in performed:
        code = str(item.get("CPT/HCPCS Code") or item.get("CPT_HCPCS") or "").strip()
        normalized_code = re.sub(r"[\s.-]", "", code).upper()
        if normalized_code and normalized_code not in deduplicated_performed:
            deduplicated_performed[normalized_code] = item
    performed = list(deduplicated_performed.values())

    encounter_type = str(request.get("encounter_type") or "OPD")
    priced_lines = price_cpt_bill_lines(performed, encounter_type=encounter_type)
    priced_count = sum(line.get("unit_charge_usd") is not None for line in priced_lines)
    if not priced_count:
        raise ValueError("None of the requested performed CPT/HCPCS codes has a valid chargemaster price for this encounter type; no draft bill was created.")

    with _STORE_LOCK:
        store = _read_store()
        _find_patient(store, patient_id)
        encounter_id = request.get("encounter_id")
        encounter_rows = _sheet_records(store, "Encounters")
        if encounter_id:
            linked = next((row for row in encounter_rows if str(_get_value(row, "Encounter_ID")) == str(encounter_id)), None)
            if linked is None or str(_get_value(linked, "Patient_ID")) != str(patient_id):
                raise ValueError("The selected encounter does not belong to this EHR patient.")
            encounter_type = str(_get_value(linked, "Type") or encounter_type)
            priced_lines = price_cpt_bill_lines(performed, encounter_type=encounter_type)
            priced_count = sum(line.get("unit_charge_usd") is not None for line in priced_lines)
            if not priced_count:
                raise ValueError("None of the requested performed CPT/HCPCS codes has a valid chargemaster price for the linked encounter type; no draft bill was created.")
        else:
            encounter_id = f"ENC-{uuid.uuid4().hex[:10].upper()}"
            encounter_sheet = store["sheets"]["Encounters"]
            _append_generated_row(store, "Encounters", encounter_sheet["headers"], {
                "Encounter_ID": encounter_id,
                "Patient_ID": str(patient_id),
                "Date": date.today().isoformat(),
                "Type": encounter_type,
                "Department": request.get("department") or "Clinical Operations",
                "Provider": request.get("provider") or "NuuCare User",
                "Facility": request.get("facility") or "NuuCare Demo",
                "Status": "Draft",
                "Primary_Diagnosis_Code": None,
                "Primary_Diagnosis": None,
                "Priority": "Routine",
                "Disposition_or_Follow_Up": "Draft billing encounter; not submitted.",
            })

        bill_id = f"BILL-{uuid.uuid4().hex[:12].upper()}"
        created_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        icd_metadata = request.get("icd10_metadata") or []
        gross_total = round(sum(float(line["gross_line_total_usd"]) for line in priced_lines if line.get("gross_line_total_usd") is not None), 2)
        unpriced = [line["cpt_hcpcs"] for line in priced_lines if line.get("unit_charge_usd") is None]
        bill_row = _append_generated_row(store, "Generated_Bills", [
            "Bill_ID", "Patient_ID", "Encounter_ID", "Encounter_Type", "ICD_10_Metadata", "CPT_HCPCS_Codes",
            "Unpriced_Codes", "Gross_Total_USD", "Status", "Created_At", "Billing_Notice",
        ], {
            "Bill_ID": bill_id,
            "Patient_ID": str(patient_id),
            "Encounter_ID": encounter_id,
            "Encounter_Type": encounter_type,
            "ICD_10_Metadata": json.dumps(icd_metadata, ensure_ascii=False),
            "CPT_HCPCS_Codes": json.dumps([line["cpt_hcpcs"] for line in priced_lines]),
            "Unpriced_Codes": json.dumps(unpriced),
            "Gross_Total_USD": gross_total,
            "Status": "Draft - Not Submitted",
            "Created_At": created_at,
            "Billing_Notice": BILLING_NOTICE,
        })
        line_rows = []
        for line in priced_lines:
            line_rows.append(_append_generated_row(store, "Generated_Bill_Lines", [
                "Bill_ID", "Patient_ID", "Encounter_ID", "Encounter_Type", "CPT_HCPCS", "Description",
                "Quantity", "Unit_Charge_USD", "Gross_Line_Total_USD", "Rate_Source_Status",
            ], {
                "Bill_ID": bill_id,
                "Patient_ID": str(patient_id),
                "Encounter_ID": encounter_id,
                "Encounter_Type": encounter_type,
                "CPT_HCPCS": line["cpt_hcpcs"],
                "Description": line["description"],
                "Quantity": line["quantity"],
                "Unit_Charge_USD": line["unit_charge_usd"],
                "Gross_Line_Total_USD": line["gross_line_total_usd"],
                "Rate_Source_Status": line["rate_source_status"],
            }))

        if "Billing_Claims" in store["sheets"]:
            claims = _sheet_records(store, "Billing_Claims")
            headers = store["sheets"]["Billing_Claims"]["headers"]
            for line in priced_lines:
                _append_generated_row(store, "Billing_Claims", headers, {
                    "Claim_ID": f"DRAFT-{bill_id}",
                    "Patient_ID": str(patient_id),
                    "Encounter_ID": encounter_id,
                    "Service_Date": date.today().isoformat(),
                    "CPT_HCPCS": line["cpt_hcpcs"],
                    "Description": line["description"],
                    "Units": line["quantity"],
                    "Charge_USD": line["gross_line_total_usd"],
                    "Claim_Status": "Draft - Not Submitted",
                    "Currency": "USD",
                })
        _write_store(store)

    return {
        **bill_row,
        "lines": line_rows,
        "gross_total_usd": gross_total,
        "unpriced_codes": unpriced,
        "icd10_metadata": icd_metadata,
        "status": "Draft - Not Submitted",
        "notice": BILLING_NOTICE,
    }
