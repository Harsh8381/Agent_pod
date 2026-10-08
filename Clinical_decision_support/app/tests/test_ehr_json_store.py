import json
from datetime import datetime

import pytest
from openpyxl import Workbook
from fastapi.testclient import TestClient

import app.services.ehr_json_service as ehr_store
from app.main import app


@pytest.fixture
def def_store_paths(tmp_path, monkeypatch):
    source = tmp_path / "Epic_Inspired_USA_50_Patient_Demo.xlsx"
    store_path = tmp_path / "Epic_Inspired_USA_50_Patient_Demo.json"
    workbook = Workbook()
    patients = workbook.active
    patients.title = "Patient_Master"
    patients.append(["Patient_ID", "MRN", "First_Name", "Last_Name", "DOB", "Email"])
    patients.append(["P100001", "MRN500001", "Olivia", "Johnson", datetime(1940, 1, 1), None])
    encounters = workbook.create_sheet("Encounters")
    encounters.append(["Encounter_ID", "Patient_ID", "Date", "Type", "Status"])
    encounters.append(["ENC001", "P100001", datetime(2026, 9, 20), "Office Visit", "Completed"])
    claims = workbook.create_sheet("Billing_Claims")
    claims.append(["Claim_ID", "Patient_ID", "Encounter_ID", "CPT_HCPCS", "Charge_USD", "Claim_Status"])
    claims.append(["CLM001", "P100001", "ENC001", "99213", 315, "Paid"])
    start = workbook.create_sheet("START_HERE")
    start.append(["Workbook title"])
    start.append(["Safety", "Synthetic data only"])
    workbook.save(source)
    monkeypatch.setattr(ehr_store, "WORKBOOK_PATH", source)
    monkeypatch.setattr(ehr_store, "JSON_PATH", store_path)
    return source, store_path


def test_excel_seed_preserves_sheets_rows_dates_and_nan_as_null(def_store_paths):
    source, store_path = def_store_paths

    payload = ehr_store.ensure_json_store()

    assert store_path.is_file()
    assert payload["sheets"]["Patient_Master"]["headers"] == [
        "Patient_ID", "MRN", "First_Name", "Last_Name", "DOB", "Email"
    ]
    patient = payload["sheets"]["Patient_Master"]["records"][0]
    assert patient["Patient_ID"] == "P100001"
    assert patient["DOB"] == "1940-01-01T00:00:00"
    assert patient["Email"] is None
    assert payload["sheets"]["START_HERE"]["raw_rows"] == [
        ["Workbook title"], ["Safety", "Synthetic data only"]
    ]
    assert json.loads(store_path.read_text(encoding="utf-8"))["source_workbook"] == source.name


def test_existing_valid_json_is_not_reseeded_from_excel(def_store_paths):
    _, store_path = def_store_paths
    ehr_store.ensure_json_store()
    existing = {
        "format": "nuucare-workbook-json-v1",
        "sheets": {"Patient_Master": {"headers": ["Patient_ID"], "records": [{"Patient_ID": "P-EDITED"}]}},
        "custom": True,
    }
    store_path.write_text(json.dumps(existing), encoding="utf-8")

    assert ehr_store.ensure_json_store() == existing
    assert "P-EDITED" in store_path.read_text(encoding="utf-8")


def test_mutations_update_patient_and_sheet_records_atomically(def_store_paths, monkeypatch):
    _, store_path = def_store_paths
    ehr_store.ensure_json_store()
    original_replace = ehr_store.os.replace
    replaced_from = []

    def track_replace(source, destination):
        replaced_from.append(source)
        assert str(source).startswith(str(store_path.parent))
        return original_replace(source, destination)

    monkeypatch.setattr(ehr_store.os, "replace", track_replace)
    ehr_store.update_patient("P100001", {"City": "Austin"})
    ehr_store.add_patient_record("P100001", "Encounters", {"Encounter_ID": "ENC002", "Type": "Follow-up"})

    saved = json.loads(store_path.read_text(encoding="utf-8"))
    patient = saved["sheets"]["Patient_Master"]["records"][0]
    assert patient["City"] == "Austin"
    assert any(row.get("Encounter_ID") == "ENC002" for row in saved["sheets"]["Encounters"]["records"])
    assert len(replaced_from) == 2


def test_patient_routes_read_create_update_record_and_delete_json_store(def_store_paths):
    client = TestClient(app)

    listed = client.get("/patients")
    assert listed.status_code == 200
    assert listed.json()["count"] == 1

    created = client.post("/patients", json={"First_Name": "Maya", "Last_Name": "Reed"})
    assert created.status_code == 200
    patient_id = created.json()["Patient_ID"]

    updated = client.put(f"/patients/{patient_id}", json={"City": "Denver"})
    assert updated.status_code == 200
    assert updated.json()["City"] == "Denver"

    chart = client.get(f"/patients/{patient_id}/chart")
    assert chart.status_code == 200
    assert chart.json()["patient"]["Patient_ID"] == patient_id

    inserted = client.post(
        f"/patients/{patient_id}/records/Encounters",
        json={"Encounter_ID": "ENC-NEW", "Type": "Office Visit"},
    )
    assert inserted.status_code == 200
    replaced = client.put(
        f"/patients/{patient_id}/records/Encounters",
        json={"records": [{"Encounter_ID": "ENC-EDIT", "Type": "Follow-up"}]},
    )
    assert replaced.status_code == 200
    assert replaced.json()[0]["Encounter_ID"] == "ENC-EDIT"

    exported = client.get("/patients/export")
    assert exported.status_code == 200
    assert "START_HERE" in exported.json()["sheets"]

    deleted = client.delete(f"/patients/{patient_id}")
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] is True


def test_draft_bill_blocks_ordered_services_and_saves_performed_lines(def_store_paths, monkeypatch):
    _, store_path = def_store_paths
    ehr_store.ensure_json_store()
    monkeypatch.setattr(
        "frontend.chargemaster.price_cpt_bill_lines",
        lambda services, encounter_type: [{
            "cpt_hcpcs": "70450",
            "description": "CT head without contrast",
            "quantity": "2",
            "unit_charge_usd": 1900.0,
            "gross_line_total_usd": 3800.0,
            "rate_source_status": "Found",
        }],
    )
    candidate = {
        "Extracted Procedure": "CT scan of head without contrast",
        "Matched Procedure/Service": "CT head without contrast",
        "CPT/HCPCS Code": "70450",
    }
    with pytest.raises(ValueError, match="explicitly performed"):
        ehr_store.generate_patient_bill("P100001", {
            "encounter_type": "OPD",
            "documentation": "CT scan of head without contrast is ordered.",
            "services": [candidate],
        })

    bill = ehr_store.generate_patient_bill("P100001", {
        "encounter_type": "OPD",
        "documentation": "CT scan of head without contrast was performed today.",
        "services": [candidate],
        "icd10_metadata": [{"ICD-10 Code": "R51.9"}],
    })

    assert bill["status"] == "Draft - Not Submitted"
    assert bill["gross_total_usd"] == 3800.0
    assert bill["lines"][0]["CPT_HCPCS"] == "70450"
    assert "R51.9" in bill["ICD_10_Metadata"]
    assert bill["Billing_Notice"].startswith("This draft contains gross chargemaster-based demo charges only.")
    saved = json.loads(store_path.read_text(encoding="utf-8"))
    assert len(saved["sheets"]["Generated_Bills"]["records"]) == 1
    assert len(saved["sheets"]["Generated_Bill_Lines"]["records"]) == 1
    assert saved["sheets"]["Billing_Claims"]["records"][-1]["Claim_Status"] == "Draft - Not Submitted"
    assert saved["sheets"]["Encounters"]["records"][-1]["Status"] == "Draft"


def test_bill_fails_clearly_when_every_requested_code_is_unpriced(def_store_paths, monkeypatch):
    monkeypatch.setattr(
        "frontend.chargemaster.price_cpt_bill_lines",
        lambda services, encounter_type: [{
            "cpt_hcpcs": "70450",
            "description": "CT head",
            "quantity": "1",
            "unit_charge_usd": None,
            "gross_line_total_usd": None,
            "rate_source_status": "Not found in chargemaster",
        }],
    )
    with pytest.raises(ValueError, match="None of the requested performed CPT/HCPCS codes has a valid chargemaster price"):
        ehr_store.generate_patient_bill("P100001", {
            "encounter_type": "OPD",
            "documentation": "CT scan of head was performed today.",
            "services": [{
                "Extracted Procedure": "CT scan of head",
                "CPT/HCPCS Code": "70450",
            }],
        })
